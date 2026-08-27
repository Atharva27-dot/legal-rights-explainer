from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt


def build_complaint_docx(report: dict, form_data: dict | None = None) -> bytes:
    """
    Creates a DOCX from the CURRENT edited complaint.
    The exporter formats the supplied text; it does not generate or
    alter legal conclusions.
    """
    report = report or {}
    form_data = form_data or {}

    complaint = str(report.get("complaint") or "").strip()
    if not complaint:
        raise ValueError("Complaint text is empty.")

    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AI-ASSISTED LEGAL COMPLAINT DRAFT")
    run.bold = True
    run.font.size = Pt(16)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Reviewable draft generated from retrieved legal provisions"
    )
    run.italic = True
    run.font.size = Pt(9)

    # User-provided case details.
    fields = [
        ("Complainant", form_data.get("name")),
        ("City / Address", form_data.get("city")),
        ("Product / Service", form_data.get("product")),
        ("Seller / Bank / Platform", form_data.get("seller")),
        ("Purchase / Transaction Date", form_data.get("purchase_date")),
        ("Amount Involved", form_data.get("amount")),
        ("Bank / Payment Platform", form_data.get("bank")),
        ("Transaction Date", form_data.get("transaction_date")),
    ]

    rows = [(label, value) for label, value in fields if str(value or "").strip()]

    if rows:
        table = document.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        for label, value in rows:
            cells = table.add_row().cells
            cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

            left = cells[0].paragraphs[0]
            left.add_run(label).bold = True

            right = cells[1].paragraphs[0]
            right.add_run(str(value))

        document.add_paragraph()

    grounding = report.get("grounding") or {}
    grounding_status = (
        grounding.get("status")
        or (report.get("case_analysis") or {}).get("grounding_status")
        or "Not available"
    )

    p = document.add_paragraph()
    p.add_run("Grounding Status: ").bold = True
    p.add_run(str(grounding_status))

    document.add_paragraph()

    known_headings = {
        "COMPLAINT / DRAFT COMPLAINT",
        "PARTIES",
        "FACTS OF THE CASE",
        "LEGAL BASIS",
        "GROUNDS",
        "RELIEF / PRAYER",
        "PRAYER",
        "DOCUMENTS / EVIDENCE",
        "DOCUMENTS",
        "DECLARATION",
    }

    for raw_line in complaint.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw_line.strip()

        if not line:
            document.add_paragraph()
            continue

        normalized = line.rstrip(":").strip().upper()

        if normalized in known_headings:
            p = document.add_paragraph()
            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(11)
        else:
            p = document.add_paragraph(line)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15

    notice = document.add_paragraph()
    notice.paragraph_format.space_before = Pt(12)
    run = notice.add_run(
        "Important: This document is an AI-assisted draft for review. "
        "It is not a substitute for professional legal advice and should "
        "be checked before submission."
    )
    run.italic = True
    run.font.size = Pt(8.5)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
